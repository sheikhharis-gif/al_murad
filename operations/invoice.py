"""Invoicing > Generate Invoice / Invoices Status.

Pick a client's trips over a date range, tick which ones to bill, choose (and
drag into order) the columns the Trip Details page should carry, optionally
work out sales tax, and download the Sales Tax Invoice as a PDF or an Excel
file - page 1 the invoice itself, page 2 the trip details. Every one
generated is logged as a GeneratedInvoice (client, company, trips, columns in
order, tax settings, totals) so Invoices Status can list them, track a status
on each, and rebuild the same invoice again on demand - the file itself is
never stored."""
import io
import re
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from urllib.parse import urlencode
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, Frame, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle,
)

from masters.models import Client, Company, TaxSettings
from .models import GeneratedInvoice, Trip

TRIP_RELATED = ("route__origin", "route__destination", "vehicle__vehicle_type", "vehicle_type",
                "sub_category", "stopover_city")

# (key, label, value-from-trip, is a money column that gets totalled)
COLUMNS = [
    ("trip_no", "Trip #", lambda t: t.trip_no, False),
    ("trip_date", "Date", lambda t: t.trip_date, False),
    ("bilty_number", "Bilty #", lambda t: t.bilty_number, False),
    ("vehicle", "Vehicle #", lambda t: t.vehicle.vehicle_number, False),
    ("vehicle_type", "Vehicle Type", lambda t: str(t.vehicle_type or t.vehicle.vehicle_type or ""), False),
    ("sub_category", "Sub-Category", lambda t: t.sub_category.name if t.sub_category_id else "", False),
    ("route", "Route", lambda t: t.route.route_code, False),
    ("weight", "Weight (Tons)", lambda t: t.weight, False),
    ("stopover_city", "Stopover City", lambda t: t.stopover_city.name if t.stopover_city_id else "", False),
    ("stopover_charges", "Stopover Charges", lambda t: t.stopover_charges, True),
    ("trip_charges", "Trip Charges", lambda t: t.freight - (t.additional_charges or 0) - (t.stopover_charges or 0), True),
    ("additional_charges", "Additional Charges", lambda t: t.additional_charges, True),
    ("total_freight", "Total Freight", lambda t: t.freight, True),
    ("remarks", "Remarks", lambda t: t.remarks, False),
]
DEFAULT_COLUMNS = ["trip_no", "trip_date", "bilty_number", "vehicle", "vehicle_type", "route", "weight",
                    "trip_charges", "additional_charges", "total_freight"]

# Full tax mode's jurisdictions (same keys as City.province) and the label
# each carries on the invoice, in the order the invoice lists them.
TAX_PROVINCES = [
    ("SINDH", "Sindh"), ("PUNJAB", "Punjab"), ("ICT", "ICT (Islamabad Capital Territory)"),
    ("KPK", "KPK (Khyber Pakhtoon Khuwa)"), ("BALOCHISTAN", "Balochistan"),
]

BLUE = "1F4E79"


def _client_columns(client):
    """A client without Sub-Categories / Stopover Charges never sees those
    two ticks - there's nothing on their trips to invoice for."""
    cols = COLUMNS
    if not client.has_sub_categories:
        cols = [c for c in cols if c[0] != "sub_category"]
    if not client.has_stopover:
        cols = [c for c in cols if c[0] not in ("stopover_city", "stopover_charges")]
    return cols


def _d(value):
    try:
        return Decimal(str(value or 0))
    except InvalidOperation:
        return Decimal(0)


def _pct(rate):
    return format(_d(rate).normalize(), "f") + "%"


def _parse_date(text):
    try:
        return date.fromisoformat(text) if text else None
    except ValueError:
        return None


@login_required
def invoice_select(request):
    client_id = request.GET.get("client") or ""
    start_date = request.GET.get("start_date") or ""
    end_date = request.GET.get("end_date") or ""
    client = Client.objects.filter(pk=client_id).first() if client_id else None
    trips = []
    if client:
        trips = Trip.objects.filter(client=client).select_related(*TRIP_RELATED)
        if start_date:
            trips = trips.filter(trip_date__gte=start_date)
        if end_date:
            trips = trips.filter(trip_date__lte=end_date)
        trips = list(trips.order_by("trip_date", "id"))

    columns = _client_columns(client) if client else []
    trip_rows = []
    for t in trips:
        cells = []
        for key, _, fn, money in columns:
            value = fn(t)
            if money:
                cells.append((key, f"{_d(value):,.0f}", True))
            elif hasattr(value, "strftime"):
                cells.append((key, value.strftime("%d-%b-%y"), False))
            else:
                cells.append((key, str(value) if value not in (None, "") else "--", False))
        trip_rows.append({"id": t.id, "total": t.freight, "cells": cells})

    return render(request, "operations/invoice_select.html", {
        "clients": Client.objects.order_by("name"),
        "companies": Company.objects.filter(is_active=True).order_by("name"),
        "client": client,
        "start_date": start_date, "end_date": end_date,
        "trip_rows": trip_rows,
        "columns": columns,
        "default_columns": DEFAULT_COLUMNS,
        "tax_provinces": TAX_PROVINCES,
        "tax_settings": TaxSettings.current(),
    })


def _tax_inputs(post):
    """Tax card fields: On/Off, Full/Partial and the 5 jurisdiction rate
    boxes (used by both modes). The boxes come pre-filled from the Tax menu
    but are editable per invoice, so what's typed here wins; a blank box falls
    back to the saved default. The rates that applied are returned as a
    snapshot to store on the invoice, so a redownload never changes if the Tax
    menu is edited later."""
    enabled = post.get("tax_enabled") == "on"
    mode = post.get("tax_mode") if post.get("tax_mode") in ("FULL", "PARTIAL") else ""
    if not enabled or not mode:
        return False, "", {}
    cfg = TaxSettings.current()

    def rate(field, default):
        raw = (post.get(field) or "").strip()
        if raw == "":
            return _d(default)
        return min(max(_d(raw), Decimal(0)), Decimal(100))

    rates = {code: rate(f"tax_{code.lower()}", getattr(cfg, f"{code.lower()}_percent")) for code, _ in TAX_PROVINCES}
    if mode == "PARTIAL" and ((post.get("tax_origin") or "").strip() or (post.get("tax_destination") or "").strip()):
        # Origin % / Destination % typed (a blank one counts as 0): one flat rate for
        # every trip's origin half / destination half instead of per-province.
        rates["origin"] = rate("tax_origin", 0)
        rates["destination"] = rate("tax_destination", 0)
    return True, mode, {k: str(v) for k, v in rates.items()}


def _compute_tax(trips, tax_enabled, tax_mode, tax_rates):
    """Returns (subtotal, tax_amount, breakdown) where breakdown is a list of
    (label, rate%, taxable base, tax amount) rows for the invoice's tax table.

    Full: each trip's whole amount is taxed at its route's ORIGIN city's
    jurisdiction rate. Partial: each trip is split 50/50 - the origin half at
    the origin city's jurisdiction rate, the destination half at the
    destination city's - worked out trip by trip, so one invoice can mix
    routes (e.g. ICT and Punjab). A city with no province set pays no tax.
    Partial with Origin % / Destination % typed on the invoice uses those two
    flat rates for every trip instead of the per-province ones."""
    subtotal = sum((_d(t.freight) for t in trips), Decimal(0))
    if not tax_enabled:
        return subtotal, Decimal(0), []

    rates = {k: _d(v) for k, v in tax_rates.items()}
    if tax_mode == "PARTIAL" and ("origin" in rates or "destination" in rates):
        origin_base = dest_base = Decimal(0)
        for t in trips:
            origin_half = (_d(t.freight) / 2).quantize(Decimal("0.01"))
            origin_base += origin_half
            dest_base += _d(t.freight) - origin_half
        origin_rate, dest_rate = rates.get("origin", Decimal(0)), rates.get("destination", Decimal(0))
        origin_tax = (origin_base * origin_rate / 100).quantize(Decimal("0.01"))
        dest_tax = (dest_base * dest_rate / 100).quantize(Decimal("0.01"))
        return subtotal, origin_tax + dest_tax, [
            ("Origin", origin_rate, origin_base, origin_tax),
            ("Destination", dest_rate, dest_base, dest_tax),
        ]

    base_by_code = {code: Decimal(0) for code, _ in TAX_PROVINCES}
    unmapped = Decimal(0)
    for t in trips:
        amount = _d(t.freight)
        if tax_mode == "FULL":
            parts = [(t.route.origin.province, amount)]
        else:
            origin_half = (amount / 2).quantize(Decimal("0.01"))
            parts = [(t.route.origin.province, origin_half), (t.route.destination.province, amount - origin_half)]
        for province, part in parts:
            if province in base_by_code:
                base_by_code[province] += part
            else:
                unmapped += part

    breakdown, tax_amount = [], Decimal(0)
    for code, label in TAX_PROVINCES:
        base, rate = base_by_code[code], rates.get(code, Decimal(0))
        amount = (base * rate / 100).quantize(Decimal("0.01"))
        tax_amount += amount
        breakdown.append((label, rate, base, amount))
    if unmapped:
        breakdown.append(("No Province Set", Decimal(0), unmapped, Decimal(0)))
    return subtotal, tax_amount, breakdown


_ONES = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve",
         "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen"]
_TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def _words_below_1000(n):
    parts = []
    if n >= 100:
        parts.append(f"{_ONES[n // 100]} Hundred")
        n %= 100
    if n >= 20:
        parts.append(_TENS[n // 10] + (f" {_ONES[n % 10]}" if n % 10 else ""))
    elif n:
        parts.append(_ONES[n])
    return " ".join(parts)


def amount_in_words(amount):
    """'Rupees Four Hundred Seventy Six Thousand Two Hundred Ninety Eight Only'
    for the whole-rupee amount printed on the invoice (thousand / million
    grouping, matching the digits shown next to it)."""
    n = int(_d(amount).quantize(Decimal("1"), rounding="ROUND_HALF_UP"))
    if n == 0:
        return "Rupees Zero Only"
    parts = []
    for size, name in ((10**9, "Billion"), (10**6, "Million"), (1000, "Thousand"), (1, "")):
        chunk, n = divmod(n, size)
        if chunk:
            parts.append(f"{_words_below_1000(chunk)} {name}".strip())
    return "Rupees " + " ".join(parts) + " Only"


def _billing_period(start, end):
    if start.year == end.year and start.month == end.month:
        return f"{start:%d}–{end:%d %b %Y}"
    return f"{start:%d %b %Y} – {end:%d %b %Y}"


def _invoice_data(invoice, trips, cols):
    """Everything both the PDF and the Excel need, worked out once."""
    cfg = TaxSettings.current()
    inv_date = invoice.created_at.astimezone(ZoneInfo("Asia/Karachi")).date()
    start = invoice.period_start or (min(t.trip_date for t in trips) if trips else inv_date)
    end = invoice.period_end or (max(t.trip_date for t in trips) if trips else inv_date)
    subtotal, tax_amount, breakdown = _compute_tax(trips, invoice.tax_enabled, invoice.tax_mode, invoice.tax_rates)

    money_idx = [i for i, c in enumerate(cols) if c[3]]
    rows, totals = [], [Decimal(0)] * len(cols)
    for t in trips:
        row = []
        for i, (_, _, fn, money) in enumerate(cols):
            value = fn(t)
            if money:
                value = _d(value)
                totals[i] += value
            row.append(value)
        rows.append(row)

    company = invoice.company
    return {
        "invoice_no": invoice.invoice_no,
        "invoice_date": inv_date,
        "due_date": inv_date + timedelta(days=invoice.payment_days),
        "payment_days": invoice.payment_days,
        "period": _billing_period(start, end),
        "bill_to": {"name": company.name, "address": company.address, "ntn": company.ntn, "strn": company.stn},
        "provider": {"name": cfg.provider_name, "address": cfg.provider_address,
                     "ntn": cfg.provider_ntn, "strn": cfg.provider_strn},
        "in_words": amount_in_words(subtotal + tax_amount),
        "notes": invoice.notes.strip() or f"Payment is due within {invoice.payment_days} days of the invoice date.",
        "subtotal": subtotal, "tax_enabled": invoice.tax_enabled, "tax_mode": invoice.tax_mode,
        "tax_amount": tax_amount, "breakdown": breakdown, "grand_total": subtotal + tax_amount,
        "headers": [c[1] for c in cols], "money_idx": money_idx, "money": [c[3] for c in cols],
        "keys": [c[0] for c in cols], "rows": rows, "totals": totals,
    }


def _safe_filename(text):
    return re.sub(r"[^A-Za-z0-9 ._()#-]", "", text).strip(" .") or "Invoice"


def _respond(invoice, trips, cols, fmt):
    data = _invoice_data(invoice, trips, cols)
    name = _safe_filename(f"Invoice {invoice.invoice_no} - {invoice.company.name}")
    if fmt == "xlsx":
        response = HttpResponse(
            _build_xlsx(data),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    else:
        response = HttpResponse(_build_pdf(data), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
    return response


@login_required
def invoice_generate_pdf(request):
    """Creates the invoice record and returns it as a PDF (format=pdf, the
    default) or an Excel file (format=xlsx)."""
    if request.method != "POST":
        return redirect("invoice_select")

    client = get_object_or_404(Client, pk=request.POST.get("client"))
    company_id = request.POST.get("company") or ""
    company = Company.objects.filter(pk=company_id).first() if company_id.isdigit() else None
    trip_ids = request.POST.getlist("trip_ids")
    problem = ("Please choose a Service Recipient (add one under Invoicing > Add Company if the list is empty)."
               if not company else "Please tick at least one trip." if not trip_ids else "")
    if problem:
        messages.error(request, problem)
        back = {"client": client.pk, "start_date": request.POST.get("start_date") or "",
                "end_date": request.POST.get("end_date") or ""}
        return redirect(f"{reverse('invoice_select')}?{urlencode(back)}")
    by_key = {c[0]: c for c in _client_columns(client)}
    # The columns arrive in the order they were dragged into on the page.
    col_keys = []
    for k in request.POST.getlist("columns"):
        if k in by_key and k not in col_keys:
            col_keys.append(k)
    if not col_keys:
        col_keys = [k for k in DEFAULT_COLUMNS if k in by_key]
    cols = [by_key[k] for k in col_keys]

    trips = list(Trip.objects.filter(pk__in=trip_ids, client=client)
                 .select_related(*TRIP_RELATED).order_by("trip_date", "id"))

    tax_enabled, tax_mode, tax_rates = _tax_inputs(request.POST)
    subtotal, tax_amount, _ = _compute_tax(trips, tax_enabled, tax_mode, tax_rates)
    try:
        payment_days = min(max(int(request.POST.get("payment_days")), 0), 365)
    except (TypeError, ValueError):
        payment_days = TaxSettings.current().payment_terms_days

    invoice = GeneratedInvoice.objects.create(
        client=client, company=company, columns=col_keys,
        period_start=_parse_date(request.POST.get("start_date")),
        period_end=_parse_date(request.POST.get("end_date")), payment_days=payment_days,
        notes=(request.POST.get("notes") or "").strip(),
        tax_enabled=tax_enabled, tax_mode=tax_mode, tax_rates=tax_rates,
        subtotal=subtotal, tax_amount=tax_amount, grand_total=subtotal + tax_amount,
        created_by=request.user if request.user.is_authenticated else None,
    )
    invoice.trips.set(trips)
    return _respond(invoice, trips, cols, request.POST.get("format"))


@login_required
def invoice_redownload(request, invoice_id):
    invoice = get_object_or_404(GeneratedInvoice.objects.select_related("client", "company"), pk=invoice_id)
    trips = list(invoice.trips.select_related(*TRIP_RELATED).order_by("trip_date", "id"))
    by_key = {c[0]: c for c in _client_columns(invoice.client)}
    cols = [by_key[k] for k in invoice.columns if k in by_key]
    return _respond(invoice, trips, cols, request.GET.get("format"))


@login_required
def invoice_status(request):
    invoices = GeneratedInvoice.objects.select_related("client", "company").order_by("-created_at")
    return render(request, "operations/invoice_status.html", {
        "invoices": invoices,
        "status_choices": GeneratedInvoice.STATUS_CHOICES,
    })


@login_required
def invoice_status_update(request, invoice_id):
    invoice = get_object_or_404(GeneratedInvoice, pk=invoice_id)
    if request.method == "POST":
        status = request.POST.get("status")
        if status in dict(GeneratedInvoice.STATUS_CHOICES):
            invoice.status = status
            invoice.save(update_fields=["status"])
    return redirect("invoice_status")


# ---------------------------------------------------------------- formatting
def _money2(value):
    return f"{value:,.2f}" if value else "-"


def _cell_text(value, money):
    if money:
        return _money2(value)
    if hasattr(value, "strftime"):
        return value.strftime("%d-%b-%y")
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    return str(value) if value not in (None, "") else ""


# ----------------------------------------------------------------------- PDF
def _build_pdf(data):
    buffer = io.BytesIO()
    doc = BaseDocTemplate(buffer, pagesize=A4, title=f"Invoice {data['invoice_no']}")
    portrait = (A4[0] - 30 * mm, A4[1] - 30 * mm)
    land = landscape(A4)
    page1_w = portrait[0]
    page2_w = land[0] - 20 * mm

    def footer(canvas, doc_):
        # Provider + invoice number on the left, page number on the right.
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        width = canvas._pagesize[0]
        canvas.drawString(12 * mm, 5 * mm, f"{data['provider']['name']}  |  Invoice {data['invoice_no']}")
        canvas.drawRightString(width - 12 * mm, 5 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.addPageTemplates([
        PageTemplate(id="P", pagesize=A4, onPage=footer, frames=[Frame(15 * mm, 15 * mm, *portrait, id="p")]),
        PageTemplate(id="L", pagesize=land, onPage=footer,
                     frames=[Frame(10 * mm, 10 * mm, page2_w, land[1] - 20 * mm, id="l")]),
    ])

    base = getSampleStyleSheet()["Normal"]
    txt = ParagraphStyle("txt", parent=base, fontSize=8.5, leading=11)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=9)
    white = ParagraphStyle("white", parent=base, fontName="Helvetica-Bold", fontSize=9, textColor=colors.white)
    white_r = ParagraphStyle("white_r", parent=white, alignment=2)
    white_c = ParagraphStyle("white_c", parent=white, fontSize=7.5, leading=9, alignment=1)
    title = ParagraphStyle("title", parent=white, fontSize=16, leading=20)
    label = ParagraphStyle("label", parent=base, fontName="Helvetica-Bold", fontSize=8.5)
    blue = colors.HexColor("#" + BLUE)
    grid = colors.HexColor("#9ca3af")
    e = escape

    def bar(left, right, width):
        # Every block on a page is exactly the frame's width, so the title bar,
        # the info rows and the tables line up edge to edge.
        t = Table([[Paragraph(left, title), Paragraph(right, white_r)]], colWidths=[width - 60 * mm, 60 * mm],
                  rowHeights=[12 * mm])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), blue), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("LEFTPADDING", (0, 0), (0, 0), 8), ("RIGHTPADDING", (1, 0), (1, 0), 8)]))
        return t

    def party(p):
        lines = [f"<font size=13>{e(p['name'])}</font>"]
        if p["address"]:
            lines.append(f"<font size=7.5>{e(p['address'])}</font>")
        ids = "  ".join(x for x in (f"NTN: {e(p['ntn'])}" if p["ntn"] else "",
                                    f"STRN: {e(p['strn'])}" if p["strn"] else "") if x)
        if ids:
            lines.append(f"<i><font size=7.5>{ids}</font></i>")
        return Paragraph("<br/>".join(lines), txt)

    els = [bar("SALES TAX INVOICE", f"SALES TAX no. {e(data['provider']['strn'])}", page1_w), Spacer(1, 4 * mm)]

    info = Table([
        [Paragraph(x, label) for x in ("Invoice Date", "Billing Period", "Payment Terms", "Due Date", "Invoice #")],
        [Paragraph(x, txt) for x in (f"{data['invoice_date']:%d %b %Y}", e(data["period"]),
                                     f"{data['payment_days']} Days", f"{data['due_date']:%d %b %Y}",
                                     f"<b>{e(data['invoice_no'])}</b>")],
    ], colWidths=[34 * mm, 34 * mm, 34 * mm, 34 * mm, 44 * mm])
    info.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")), ("BOX", (0, 0), (-1, -1), 0.6, grid),
        ("LINEBELOW", (0, 0), (-1, 0), 0.4, grid), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    els += [info, Spacer(1, 5 * mm)]

    parties = Table([
        [Paragraph("BILL TO / CUSTOMER", white), Paragraph("SERVICE PROVIDER", white)],
        [party(data["bill_to"]), party(data["provider"])],
    ], colWidths=[90 * mm, 90 * mm], rowHeights=[None, 32 * mm])
    parties.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), blue), ("BOX", (0, 1), (-1, 1), 0.6, grid),
        ("LINEAFTER", (0, 1), (0, 1), 0.6, grid), ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    els += [parties, Spacer(1, 5 * mm)]

    desc = Table([
        [Paragraph("DESCRIPTION OF SERVICES", white), Paragraph("AMOUNT (PKR)", white_r)],
        [Paragraph("Transportation Services", txt), Paragraph(f"{data['subtotal']:,.0f}", ParagraphStyle("r", parent=txt, alignment=2))],
    ], colWidths=[130 * mm, 50 * mm], rowHeights=[None, 12 * mm])
    desc.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), blue), ("BOX", (0, 1), (-1, 1), 0.6, grid),
        ("LINEAFTER", (0, 1), (0, 1), 0.6, grid), ("VALIGN", (0, 1), (-1, 1), "MIDDLE"),
    ]))
    els += [desc, Spacer(1, 5 * mm)]

    right = ParagraphStyle("right", parent=txt, alignment=2)
    if data["tax_enabled"] and data["breakdown"]:
        rows = [[Paragraph("TAX JURISDICTION", white_c),
                 Paragraph("TAX RATE", white_c), Paragraph("TAXABLE AMOUNT (PKR)", white_c),
                 Paragraph("TAX AMOUNT (PKR)", white_c)]]
        for lbl, rate, base, amount in data["breakdown"]:
            rows.append([Paragraph(e(lbl), txt), Paragraph(_pct(rate), right),
                         Paragraph(f"{base:,.0f}" if base else "-", right),
                         Paragraph(f"{amount:,.0f}" if amount else "-", right)])
        tax = Table(rows, colWidths=[72 * mm, 32 * mm, 38 * mm, 38 * mm])
        tax.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), blue), ("GRID", (0, 0), (-1, -1), 0.4, grid),
                                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        els += [tax, Spacer(1, 5 * mm)]

    total_label = "TOTAL INVOICE AMOUNT (INCL. SALES TAX)" if data["tax_enabled"] else "TOTAL INVOICE AMOUNT"
    total = Table([[Paragraph(f"<b>{total_label}</b>", ParagraphStyle("tl", parent=txt, fontSize=10)),
                    Paragraph(f"<b>{data['grand_total']:,.0f}</b>", ParagraphStyle("tr", parent=txt, fontSize=11, alignment=2))]],
                  colWidths=[130 * mm, 50 * mm], rowHeights=[12 * mm])
    total.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#eaf0f7")),
                               ("LINEABOVE", (0, 0), (-1, 0), 1.2, blue), ("LINEBELOW", (0, 0), (-1, 0), 1.2, blue),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    notes = Table([
        [Paragraph("NOTES", white)],
        [Paragraph(e(data["notes"]).replace("\n", "<br/>"), txt)],
    ], colWidths=[180 * mm])
    notes.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), blue), ("BOX", (0, 1), (-1, 1), 0.6, grid),
        ("TOPPADDING", (0, 1), (-1, 1), 6), ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
    ]))
    words = Table([[Paragraph(f"<b>Amount in words:</b> {e(data['in_words'])}", txt)]], colWidths=[180 * mm])
    words.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, grid), ("TOPPADDING", (0, 0), (-1, -1), 6),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    els += [total, Spacer(1, 2 * mm), words, Spacer(1, 5 * mm), notes]

    # ---- page 2: trip details (landscape)
    els += [NextPageTemplate("L"), PageBreak(), bar("TRIP DETAILS", f"Invoice # {e(data['invoice_no'])}", page2_w),
            Spacer(1, 4 * mm)]
    meta = Table([
        [Paragraph(x, label) for x in ("Invoice #", "Billing Period", "Customer", "Service Provider")],
        [Paragraph(e(x), txt) for x in (data["invoice_no"], data["period"], data["bill_to"]["name"], data["provider"]["name"])],
    ], colWidths=[page2_w / 4] * 4)
    meta.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, grid), ("INNERGRID", (0, 0), (-1, -1), 0.4, grid),
                              ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                              ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    els += [meta, Spacer(1, 4 * mm)]

    head = [Paragraph("S.no", white_c)] + [Paragraph(e(h), white_c) for h in data["headers"]]
    body = [head]
    for n, row in enumerate(data["rows"], start=1):
        cells = [Paragraph(str(n), ParagraphStyle("sn", parent=small, alignment=2))]
        for value, money in zip(row, data["money"]):
            text = _cell_text(value, money)
            cells.append(Paragraph(e(text), ParagraphStyle("c", parent=small, alignment=2 if money or isinstance(value, Decimal) else 0)))
        body.append(cells)
    foot = [Paragraph("<b>TOTAL</b>", small)] + [""] * len(data["headers"])
    for i in data["money_idx"]:
        foot[i + 1] = Paragraph(f"<b>{_money2(data['totals'][i])}</b>", ParagraphStyle("f", parent=small, alignment=2))
    body.append(foot)

    widths = [10 * mm] + [(page2_w - 10 * mm) / max(len(data["headers"]), 1)] * len(data["headers"])
    trips_table = Table(body, colWidths=widths, repeatRows=1)
    trips_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), blue), ("GRID", (0, 0), (-1, -1), 0.3, grid),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#eaf0f7")), ("LINEABOVE", (0, -1), (-1, -1), 1, blue), ("SPAN", (0, -1), (1, -1)),
    ]))
    els.append(trips_table)
    doc.build(els)
    return buffer.getvalue()


# --------------------------------------------------------------------- Excel
def _build_xlsx(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    thin = Side(style="thin", color="9CA3AF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    blue_fill = PatternFill("solid", fgColor=BLUE)
    total_fill = PatternFill("solid", fgColor="EAF0F7")
    head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)

    def put(sheet, rng, value=None, font=None, fill=None, align=None, border=None, fmt=None):
        first = rng.split(":")[0]
        if ":" in rng:
            sheet.merge_cells(rng)
        cells = sheet[rng] if ":" in rng else ((sheet[rng],),)
        for row in cells:
            for c in row:
                if fill:
                    c.fill = fill
                if border:
                    c.border = border
        c = sheet[first]
        c.value = value
        if font:
            c.font = font
        if align:
            c.alignment = align
        if fmt:
            c.number_format = fmt

    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lab = Font(name="Calibri", bold=True, size=10)
    val = Font(name="Calibri", size=10)

    for i in range(1, 11):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.sheet_view.showGridLines = False

    put(ws, "A1:H1", "SALES TAX INVOICE", Font(name="Calibri", bold=True, color="FFFFFF", size=16), blue_fill,
        Alignment(vertical="center"))
    put(ws, "I1:J1", f"SALES TAX no. {data['provider']['strn']}", Font(name="Calibri", bold=True, color="FFFFFF", size=9),
        blue_fill, Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[1].height = 26

    for rng, text in (("A3:B3", "Invoice Date"), ("C3:D3", "Billing Period"), ("E3:F3", "Payment Terms"),
                      ("G3:H3", "Due Date"), ("I3:J3", "Invoice #")):
        put(ws, rng, text, lab, align=left)
    for rng, text in (("A4:B4", f"{data['invoice_date']:%d %b %Y}"), ("C4:D4", data["period"]),
                      ("E4:F4", f"{data['payment_days']} Days"), ("G4:H4", f"{data['due_date']:%d %b %Y}"),
                      ("I4:J4", data["invoice_no"])):
        put(ws, rng, text, val, align=left, border=Border(bottom=thin))

    put(ws, "A6:D6", "BILL TO / CUSTOMER", head_font, blue_fill, left)
    put(ws, "E6:J6", "SERVICE PROVIDER", head_font, blue_fill, left)
    for (c1, c2), p in ((("A", "D"), data["bill_to"]), (("E", "J"), data["provider"])):
        put(ws, f"{c1}7:{c2}7", p["name"], Font(name="Calibri", size=14), align=left, border=box)
        put(ws, f"{c1}8:{c2}9", p["address"], Font(name="Calibri", size=8), align=top_left, border=box)
        ids = "   ".join(x for x in (f"NTN: {p['ntn']}" if p["ntn"] else "", f"STRN: {p['strn']}" if p["strn"] else "") if x)
        put(ws, f"{c1}10:{c2}10", ids, Font(name="Calibri", size=8, italic=True), align=left, border=box)
    ws.row_dimensions[7].height = 24
    ws.row_dimensions[8].height = 18
    ws.row_dimensions[9].height = 18

    put(ws, "A13:H13", "DESCRIPTION OF SERVICES", head_font, blue_fill, left)
    put(ws, "I13:J13", "AMOUNT (PKR)", head_font, blue_fill, center)
    put(ws, "A14:H14", "Transportation Services", val, align=left, border=box)
    put(ws, "I14:J14", float(data["subtotal"]), val, align=right, border=box, fmt="#,##0")
    ws.row_dimensions[14].height = 30

    r = 16
    if data["tax_enabled"] and data["breakdown"]:
        put(ws, f"A{r+1}:D{r+1}", "TAX JURISDICTION", head_font, blue_fill, center)
        put(ws, f"E{r+1}:F{r+1}", "TAX RATE", head_font, blue_fill, center)
        put(ws, f"G{r+1}:H{r+1}", "TAXABLE AMOUNT (PKR)", head_font, blue_fill, center)
        put(ws, f"I{r+1}:J{r+1}", "TAX AMOUNT (PKR)", head_font, blue_fill, center)
        r += 2
        for lbl, rate, base, amount in data["breakdown"]:
            put(ws, f"A{r}:D{r}", lbl, val, align=left, border=box)
            put(ws, f"E{r}:F{r}", _pct(rate), val, align=right, border=box)
            put(ws, f"G{r}:H{r}", float(base), val, align=right, border=box, fmt='#,##0;-#,##0;"-"')
            put(ws, f"I{r}:J{r}", float(amount), val, align=right, border=box, fmt='#,##0;-#,##0;"-"')
            r += 1
        r += 1
    r += 1
    heavy = Border(top=Side(style="medium", color=BLUE), bottom=Side(style="medium", color=BLUE))
    put(ws, f"A{r}:H{r}", "TOTAL INVOICE AMOUNT (INCL. SALES TAX)" if data["tax_enabled"] else "TOTAL INVOICE AMOUNT",
        Font(name="Calibri", bold=True, size=11), total_fill, left, heavy)
    put(ws, f"I{r}:J{r}", float(data["grand_total"]), Font(name="Calibri", bold=True, size=12), total_fill, right, heavy,
        fmt="#,##0")
    ws.row_dimensions[r].height = 30
    put(ws, f"A{r+1}:J{r+1}", f"Amount in words: {data['in_words']}", Font(name="Calibri", bold=True, size=10),
        align=left, border=box)
    ws.row_dimensions[r + 1].height = 24
    put(ws, f"A{r+3}:J{r+3}", "NOTES", head_font, blue_fill, left)
    put(ws, f"A{r+4}:J{r+4}", data["notes"], val, align=top_left, border=box)
    ws.row_dimensions[r + 4].height = max(45, 15 * (data["notes"].count("\n") + 1 + len(data["notes"]) // 110))

    # ---- Trip Details
    td = wb.create_sheet("Trip Details")
    td.sheet_view.showGridLines = False
    headers = ["S.no"] + data["headers"]
    n = len(headers)
    last = get_column_letter(n)
    put(td, f"A1:{last}1", "TRIP DETAILS", Font(name="Calibri", bold=True, color="FFFFFF", size=16), blue_fill,
        Alignment(vertical="center"))
    td.row_dimensions[1].height = 26

    meta = [("Invoice #", data["invoice_no"]), ("Billing Period", data["period"]),
            ("Customer", data["bill_to"]["name"]), ("Service Provider", data["provider"]["name"])]
    groups = min(4, n)
    size, extra = divmod(n, groups)
    start = 1
    for gi in range(groups):
        end = start + size + (1 if gi < extra else 0) - 1
        a, b = get_column_letter(start), get_column_letter(end)
        put(td, f"{a}3:{b}3", meta[gi][0], lab, PatternFill("solid", fgColor="F3F4F6"), left, box)
        put(td, f"{a}4:{b}4", meta[gi][1], val, align=left, border=box)
        start = end + 1

    hr = 6
    for i, h in enumerate(headers, start=1):
        c = td.cell(row=hr, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = head_font, blue_fill, center, box
    td.row_dimensions[hr].height = 32

    money_fmt = '#,##0.00;-#,##0.00;"-"'
    r = hr + 1
    for sno, row in enumerate(data["rows"], start=1):
        td.cell(row=r, column=1, value=sno).font = val
        for i, (value, money) in enumerate(zip(row, data["money"]), start=2):
            c = td.cell(row=r, column=i, value=float(value) if isinstance(value, Decimal) else (value if value != "" else None))
            c.font, c.border = val, box
            if money:
                c.number_format = money_fmt
            elif isinstance(value, Decimal):
                c.number_format = "0.00"
            elif hasattr(value, "strftime"):
                c.number_format = "d-mmm-yy"
                c.alignment = Alignment(horizontal="right")
        td.cell(row=r, column=1).border = box
        r += 1
    put(td, f"A{r+1}", "TOTAL", Font(name="Calibri", bold=True, size=11), total_fill, left, heavy)
    for i in range(2, n + 1):
        c = td.cell(row=r + 1, column=i)
        c.fill, c.border = total_fill, heavy
        if (i - 2) in data["money_idx"]:
            c.value = float(data["totals"][i - 2])
            c.font = Font(name="Calibri", bold=True, size=11)
            c.number_format = money_fmt
    td.column_dimensions["A"].width = 7
    for i, key in enumerate(["sno"] + data["keys"], start=1):
        if i > 1:
            td.column_dimensions[get_column_letter(i)].width = 30 if key == "remarks" else (12 if key in ("route", "trip_no") else 16)
    td.freeze_panes = td.cell(row=hr + 1, column=1)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
