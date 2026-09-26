"""Excel versions of the two invoice templates ("invoice template (2).xlsx"),
reproduced cell for cell - Segoe UI throughout, the 1F4E78 blue bars, the
A6A6A6 thin borders, the same merges, row heights, number formats, formulas
and print setup.

  * Tax invoice     -> sheets "Invoice" + "Trips Summary"
  * Non-tax invoice -> one sheet "NON-TAX INVOICE"

The tables are the columns ticked (and dragged into order) on Generate
Invoice, so their width varies; everything else follows the template.
"""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

FONT = "Segoe UI"
BLUE, TEXT, LIGHT, GREY, FOOT, WHITE = "FF1F4E78", "FF1F2937", "FFF3F6F9", "FFA6A6A6", "FF666666", "FFFFFFFF"
THIN = Side(style="thin", color=GREY)
MEDIUM = Side(style="medium", color=BLUE)
ACCOUNTING0 = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
ACCOUNTING2 = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
DATE_FMT = "[$-409]d\\-mmm\\-yy;@"


def _font(size, bold=False, italic=False, color=TEXT):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _put(ws, rng, value=None, font=None, fill=None, align=None, nf=None, left=None, right=None, top=None, bottom=None):
    """Merge `rng` (if it spans cells), write `value` in its first cell and give
    every cell the fill; borders go on the outer edge only, like Excel does
    for a merged range (top/bottom on every row-edge cell, left/right on the
    end columns)."""
    ws.merge_cells(rng) if ":" in rng and rng.split(":")[0] != rng.split(":")[1] else None
    cells = ws[rng] if ":" in rng else ((ws[rng],),)
    nrows, ncols = len(cells), len(cells[0])
    for ri, row in enumerate(cells):
        for ci, c in enumerate(row):
            if fill:
                c.fill = fill
            c.border = Border(
                left=left if ci == 0 else None, right=right if ci == ncols - 1 else None,
                top=top if ri == 0 else None, bottom=bottom if ri == nrows - 1 else None)
    first = cells[0][0]
    first.value = value
    if font:
        first.font = font
    if align:
        first.alignment = align
    if nf:
        first.number_format = nf
    return first


def _box():
    return dict(left=THIN, right=THIN, top=THIN, bottom=THIN)


A_LEFT = Alignment(horizontal="left", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")
A_CENTER = Alignment(horizontal="center", vertical="center")
A_V = Alignment(vertical="center")
A_CELL = Alignment(horizontal="left", vertical="center", wrap_text=True)
A_HEAD = Alignment(horizontal="left", vertical="center", wrap_text=True)
A_HEAD_R = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _party(name, address, ntn, strn, name_size):
    """Rich text like the template: the name large, the rest 9 pt below it."""
    rest = [address.strip()] if address and address.strip() else []
    ids = "    ".join(x for x in (f"NTN: {ntn}" if ntn else "", f"STRN: {strn}" if strn else "") if x)
    if ids:
        rest.append(ids)
    big = TextBlock(InlineFont(rFont=FONT, sz=name_size, color=TEXT), name)
    if not rest:
        return CellRichText([big])
    return CellRichText([big, TextBlock(InlineFont(rFont=FONT, sz=9, color=TEXT), "\n" + "\n".join(rest))])


def _pct(rate):
    rate = Decimal(str(rate))
    return float(rate / 100), ("0%" if rate == rate.to_integral_value() else "0.00%")


def _split(n_cols, groups):
    """[(first, last), ...] 1-based column ranges splitting n_cols into `groups` runs."""
    groups = min(groups, n_cols)
    size, extra = divmod(n_cols, groups)
    out, start = [], 1
    for i in range(groups):
        end = start + size + (1 if i < extra else 0) - 1
        out.append((start, end))
        start = end + 1
    return out


def _page(ws, orientation, last_row, last_col):
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = 16.8
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.oddFooter.center.text = "Page &P of &N"
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"


# Column widths in the templates, by position (S.no first); a wide text column overrides.
TRIPS_SUMMARY_WIDTHS = [14, 15, 16, 15, 16, 15, 12, 15, 16, 16.89, 18.33, 15]
NON_TAX_WIDTHS = [14, 15, 16, 15, 16, 15, 13.66, 15, 16, 16.89]
_MIN_WIDTH = {"remarks": 30, "stopover_city": 16, "sub_category": 15}


def _width(widths, i, key):
    base = widths[i - 1] if i <= len(widths) else 16
    return max(base, _MIN_WIDTH.get(key, 0))


def _col_format(key, money):
    if money:
        return ACCOUNTING2
    if key == "trip_date":
        return DATE_FMT
    if key in ("weight", "route"):  # the template formats its Route column #,##0.00 too
        return "#,##0.00"
    return None


def _table(ws, data, header_row, height_rows=None, wrap_header=True):
    """Header + body of the trips table (S.no first) starting at header_row.
    Returns (first_data_row, last_data_row)."""
    headers = ["S.no"] + data["headers"]
    keys = ["sno"] + data["keys"]
    money = [False] + data["money"]
    for i, h in enumerate(headers, start=1):
        _put(ws, f"{get_column_letter(i)}{header_row}", h, _font(10, True, color=WHITE), _fill(BLUE),
             (A_HEAD_R if wrap_header else A_RIGHT) if money[i - 1] else (A_HEAD if wrap_header else A_LEFT), **_box())
    r = header_row + 1
    for n, row in enumerate(data["rows"], start=1):
        for i, (key, is_money) in enumerate(zip(keys, money), start=1):
            value = n if i == 1 else row[i - 2]
            if isinstance(value, Decimal):
                value = float(value)
            elif value == "":
                value = None
            _put(ws, f"{get_column_letter(i)}{r}", value, _font(10), None, A_CELL,
                     nf=_col_format(key, is_money), **_box())
        if height_rows:
            ws.row_dimensions[r].height = height_rows
        r += 1
    return header_row + 1, r - 1


def _total_row(ws, data, row, first_data, last_data, n_cols):
    """TOTAL bar: label merged over the columns before the first money column,
    a SUM under each money column."""
    money = [False] + data["money"]
    first_money = next((i for i, m in enumerate(money, start=1) if m), n_cols + 1)
    edge = dict(left=THIN, right=THIN, top=MEDIUM, bottom=MEDIUM)
    label_to = max(first_money - 1, 1)
    _put(ws, f"A{row}:{get_column_letter(label_to)}{row}" if label_to > 1 else f"A{row}", "TOTAL",
         _font(11, True), _fill(LIGHT), A_V, **edge)
    for i in range(label_to + 1, n_cols + 1):
        col = get_column_letter(i)
        if money[i - 1]:
            _put(ws, f"{col}{row}", f"=SUM({col}{first_data}:{col}{last_data})", _font(11, True), _fill(LIGHT), A_V,
                 nf="#,##0.00", **edge)
        else:
            _put(ws, f"{col}{row}", None, None, _fill(LIGHT), None, top=MEDIUM, bottom=MEDIUM)
    ws.row_dimensions[row].height = 19.95


def _freight_ref(data, sheet_prefix, total_row):
    """Formula pointing at the Total Freight sum, or the plain amount if that
    column isn't on the invoice."""
    if "total_freight" in data["keys"]:
        col = get_column_letter(data["keys"].index("total_freight") + 2)
        return f"={sheet_prefix}{col}{total_row}"
    return float(data["subtotal"])


def _info_row(ws, labels_row, pairs, n_cols, label_font, value_font, boxed):
    """The Invoice Date / Billing Period / ... strip, spread over n_cols."""
    for (a, b), (label, value) in zip(_split(n_cols, len(pairs)), pairs):
        ra, rb = f"{get_column_letter(a)}", f"{get_column_letter(b)}"
        rng = (lambda row: f"{ra}{row}:{rb}{row}" if b > a else f"{ra}{row}")
        if boxed:
            _put(ws, rng(labels_row), label, label_font, _fill(LIGHT), A_V, **_box())
            _put(ws, rng(labels_row + 1), value, value_font, None, A_LEFT, **_box())
        else:
            _put(ws, rng(labels_row), label, label_font, None, A_V, bottom=THIN)
            _put(ws, rng(labels_row + 1), value, value_font, None, A_V, bottom=THIN)


def _bullets(notes):
    return [f"• {line}" for line in notes]


# ------------------------------------------------------------------ tax invoice
def _tax_workbook(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ts = wb.create_sheet("Trips Summary")

    # ===== Trips Summary (page 2)
    headers = ["S.no"] + data["headers"]
    n_cols = len(headers)
    last = get_column_letter(n_cols)
    _put(ts, f"A1:{last}1", "TRIPS SUMMARY", _font(18, True, color=WHITE), _fill(BLUE), A_LEFT)
    ts.row_dimensions[1].height = 30
    meta = [("Invoice #", data["invoice_no"]), ("Billing Period", data["period"]),
            ("Customer", data["bill_to"]["name"]), ("Service Provider", data["provider"]["name"])]
    _info_row(ts, 3, meta, n_cols, _font(10, True), _font(10), boxed=True)
    ts.row_dimensions[3].height = ts.row_dimensions[4].height = 19.95
    ts.row_dimensions[6].height = 30
    first_data, last_data = _table(ts, data, 6)
    spacer = last_data + 1
    ts.row_dimensions[spacer].height = 17.4
    total_row = spacer + 1
    _total_row(ts, data, total_row, first_data, last_data, n_cols)
    footer_row = total_row + 2
    _put(ts, f"A{footer_row}:{last}{footer_row}", "Page 2 of 2", _font(8, italic=True, color=FOOT), None, Alignment(horizontal="center"))
    for i, key in enumerate(["sno"] + data["keys"], start=1):
        ts.column_dimensions[get_column_letter(i)].width = _width(TRIPS_SUMMARY_WIDTHS, i, key)
    _page(ts, "landscape", footer_row, n_cols)

    # ===== Invoice (page 1) - 10 columns of 17
    for i in range(1, 11):
        ws.column_dimensions[get_column_letter(i)].width = 17
    _put(ws, "A1:H1", "SALES TAX INVOICE", _font(18, True, color=WHITE), _fill(BLUE), A_LEFT)
    _put(ws, "I1", "SALES TAX no.", _font(8, True, color=WHITE), _fill(BLUE), A_RIGHT)
    _put(ws, "J1", data["sales_tax_no"] or None, _font(9, True, color=WHITE), _fill(BLUE), Alignment(horizontal="left", vertical="center"))
    _info_row(ws, 3, [("Invoice Date", f"{data['invoice_date']:%d %b %Y}"), ("Billing Period", data["period"]),
                      ("Payment Terms", f"{data['payment_days']} Days"), ("Due Date", f"{data['due_date']:%d %b %Y}"),
                      ("Invoice #", data["invoice_no"])], 10, _font(9, True), _font(10), boxed=False)

    _put(ws, "A6:E6", "BILL TO / CUSTOMER", _font(10, True, color=WHITE), _fill(BLUE), A_V, **_box())
    _put(ws, "F6:J6", "SERVICE PROVIDER", _font(10, True, color=WHITE), _fill(BLUE), A_V, **_box())
    top = Alignment(vertical="top", wrap_text=True)
    b, p = data["bill_to"], data["provider"]
    _put(ws, "A7:E11", _party(b["name"], b["address"], b["ntn"], b["strn"], 16), _font(9), None, top, **_box())
    _put(ws, "F7:J11", _party(p["name"], p["address"], p["ntn"], p["strn"], 16), _font(9), None, top, **_box())

    _put(ws, "A13:H13", "DESCRIPTION OF SERVICES", _font(9, True, color=WHITE), _fill(BLUE), A_V, **_box())
    _put(ws, "I13:J13", "AMOUNT (PKR)", _font(9, True, color=WHITE), _fill(BLUE), A_RIGHT, **_box())
    _put(ws, "A14:H15", "Transportation Services", _font(10), None, A_LEFT, **_box())
    _put(ws, "I14:J15", _freight_ref(data, "'Trips Summary'!", total_row), _font(10), None, A_RIGHT, nf=ACCOUNTING0, **_box())

    _put(ws, "A17:D17", "TAX JURISDICTION", _font(8, True, color=WHITE), _fill(BLUE), A_HEAD, **_box())
    _put(ws, "E17:F17", "TAX RATE", _font(8, True, color=WHITE), _fill(BLUE), Alignment(horizontal="center", vertical="center", wrap_text=True), **_box())
    _put(ws, "G17:H17", "TAXABLE AMOUNT (PKR)", _font(8, True, color=WHITE), _fill(BLUE), Alignment(horizontal="center", vertical="center", wrap_text=True), **_box())
    _put(ws, "I17:J17", "TAX AMOUNT (PKR)", _font(8, True, color=WHITE), _fill(BLUE), Alignment(horizontal="right", vertical="center", wrap_text=True), **_box())
    r = 18
    for label, rate, base, amount in data["breakdown"]:
        short = "KPK" if label.startswith("KPK") else label
        rate_value, rate_fmt = _pct(rate)
        _put(ws, f"A{r}:D{r}", short, _font(10), None, A_LEFT, **_box())
        _put(ws, f"E{r}:F{r}", rate_value, _font(10), None, A_RIGHT, nf=rate_fmt, **_box())
        _put(ws, f"G{r}:H{r}", "=I14" if base and base == data["subtotal"] else float(base), _font(10), None, A_RIGHT,
             nf=ACCOUNTING0, **_box())
        _put(ws, f"I{r}:J{r}", f"=G{r}*E{r}", _font(10), None, A_RIGHT, nf=ACCOUNTING0, **_box())
        r += 1
    last_tax = r - 1
    total_top = r + 1
    edge = dict(left=THIN, right=THIN, top=MEDIUM, bottom=MEDIUM)
    _put(ws, f"A{total_top}:H{total_top + 1}", "TOTAL INVOICE AMOUNT (INCL. SALES TAX)", _font(12, True), _fill(LIGHT), A_LEFT, **edge)
    _put(ws, f"I{total_top}:J{total_top + 1}", f"=I14+SUM(I18:J{last_tax})", _font(13, True, color=BLUE), _fill(LIGHT), A_RIGHT,
         nf=ACCOUNTING0, **edge)
    words_row = total_top + 3
    _put(ws, f"A{words_row}:B{words_row}", "AMOUNT IN WORDS", _font(9, True), None, None, bottom=THIN)
    _put(ws, f"C{words_row}:J{words_row}", data["in_words"], _font(10, italic=True), None, None, bottom=THIN)
    notes_row = words_row + 2
    _put(ws, f"A{notes_row}:J{notes_row}", "NOTES", _font(10, True, color=WHITE), _fill(BLUE), None, **_box())
    lines = _bullets(data["notes"])
    for i, line in enumerate(lines, start=1):
        _put(ws, f"A{notes_row + i}:J{notes_row + i}", line, _font(9), None, A_V, bottom=THIN)
    foot = notes_row + max(len(lines), 3) + 4
    _put(ws, f"A{foot}:J{foot}", "Page 1 of 2", _font(8, italic=True, color=FOOT), None, Alignment(horizontal="center"))
    for row in range(2, foot + 1):
        ws.row_dimensions[row].height = 19.95
    ws.row_dimensions[1].height = 30
    _page(ws, "portrait", foot, 10)
    return wb


# --------------------------------------------------------------- non-tax invoice
def _nontax_workbook(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "NON-TAX INVOICE"
    headers = ["S.no"] + data["headers"]
    keys = ["sno"] + data["keys"]
    n_cols = max(len(headers), 10)
    L = get_column_letter
    half = n_cols // 2
    for i in range(1, n_cols + 1):
        ws.column_dimensions[L(i)].width = _width(NON_TAX_WIDTHS, i, keys[i - 1] if i <= len(keys) else "")

    b, p = data["bill_to"], data["provider"]
    _put(ws, f"A1:{L(half)}1", "BILL TO / CUSTOMER", _font(14, True, color=WHITE), _fill(BLUE), A_V, **_box())
    _put(ws, f"{L(half + 1)}1:{L(n_cols)}1", "SERVICE PROVIDER", _font(14, True, color=WHITE), _fill(BLUE), A_V, **_box())
    top = Alignment(vertical="top", wrap_text=True)
    _put(ws, f"A2:{L(half)}7", _party(b["name"], b["address"], b["ntn"], b["strn"], 18), _font(9), None, top, **_box())
    _put(ws, f"{L(half + 1)}2:{L(n_cols)}7", _party(p["name"], p["address"], p["ntn"], p["strn"], 18), _font(9), None, top, **_box())
    _info_row(ws, 9, [("Invoice Date", f"{data['invoice_date']:%d %b %Y}"), ("Billing Period", data["period"]),
                      ("Payment Terms", f"{data['payment_days']} Days"), ("Due Date", f"{data['due_date']:%d %b %Y}"),
                      ("Invoice #", data["invoice_no"])], n_cols, _font(9, True), _font(10), boxed=False)
    _put(ws, f"A12:{L(n_cols)}12", "TRIPS SUMMARY", _font(16, True, color=WHITE), _fill(BLUE), A_LEFT)
    first_data, last_data = _table(ws, data, 14, height_rows=19.95, wrap_header=False)
    total_row = last_data + 2
    _total_row(ws, data, total_row, first_data, last_data, len(headers))
    # the template's table is exactly the sheet width; pad the bar if the table is narrower
    inv_row = total_row + 2
    edge = dict(left=THIN, right=THIN, top=MEDIUM, bottom=MEDIUM)
    amount_cols = 2
    _put(ws, f"A{inv_row}:{L(n_cols - amount_cols)}{inv_row}", "TOTAL INVOICE AMOUNT", _font(12, True), _fill(LIGHT), A_LEFT, **edge)
    _put(ws, f"{L(n_cols - amount_cols + 1)}{inv_row}:{L(n_cols)}{inv_row}",
         _freight_ref(data, "", total_row), _font(13, True, color=BLUE), _fill(LIGHT), A_RIGHT, nf=ACCOUNTING0, **edge)
    words_row = inv_row + 2
    _put(ws, f"A{words_row}:B{words_row}", "AMOUNT IN WORDS", _font(9, True), None, None, bottom=THIN)
    _put(ws, f"C{words_row}:{L(n_cols)}{words_row}", data["in_words"], _font(10, italic=True), None, None, bottom=THIN)
    notes_row = words_row + 2
    _put(ws, f"A{notes_row}:{L(n_cols)}{notes_row}", "NOTES", _font(10, True, color=WHITE), _fill(BLUE), None, **_box())
    # a single page has no "Page 2" to refer to
    lines = _bullets([n for n in data["notes"] if "Page 2" not in n])
    for i, line in enumerate(lines, start=1):
        _put(ws, f"A{notes_row + i}:{L(n_cols)}{notes_row + i}", line, _font(9), None, A_V, bottom=THIN)
    last = notes_row + len(lines)

    for row in range(3, last + 1):
        ws.row_dimensions[row].height = 19.95
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[12].height = ws.row_dimensions[14].height = 25.05
    ws.row_dimensions[inv_row].height = 30
    ws.row_dimensions[total_row].height = 19.95
    _page(ws, "landscape", last, n_cols)
    return wb


def build(data):
    wb = _tax_workbook(data) if data["tax_enabled"] else _nontax_workbook(data)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
