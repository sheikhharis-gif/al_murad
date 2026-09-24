"""Client Rates page - Excel and PDF downloads of a client's rate sheet
(fuel-indexed route rates + dedicated vehicle rates), same columns as on screen."""
import io
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Client, ClientRate, DedicatedRate

RATE_HEADERS = [
    "Route", "Fuel Product", "Current Fuel Price", "Current Rate", "Effective %",
    "Rate Subject to Revision", "Updated Fuel Price", "Fuel Price Change %",
    "Rate Adjustment", "Updated Trip Cost", "Weight (Tons)", "Vehicle Type", "Effective Date",
]
DEDICATED_HEADERS = [
    "Vehicle ID", "Fixed Cost", "Month", "Fuel Avg", "Fuel Price", "Variable Cost", "Route",
    "Distance Source", "Distance (Km)", "Weight (Tons)", "Vehicle Type", "Effective Date",
]
# Values that are percentages (shown with a % sign)
PERCENT_COLS = {"Effective %", "Fuel Price Change %"}


def _date_range(request):
    """From / To effective-date filter from the download form (either may be blank)."""
    def parse(name):
        try:
            return date.fromisoformat(request.GET.get(name) or "")
        except ValueError:
            return None
    return parse("from"), parse("to")


def _in_range(queryset, start, end):
    if start:
        queryset = queryset.filter(effective_date__gte=start)
    if end:
        queryset = queryset.filter(effective_date__lte=end)
    return queryset


def _period_label(start, end):
    if not start and not end:
        return "All dates"
    return f"Effective {start:%d-%b-%Y} to {end:%d-%b-%Y}" if start and end else (
        f"Effective from {start:%d-%b-%Y}" if start else f"Effective up to {end:%d-%b-%Y}")


def _rate_headers(client):
    """Sub-Category / Rate Type columns only for clients that use them."""
    return (["Sub-Category", "Rate Type"] if client.has_sub_categories else []) + RATE_HEADERS


def _rate_rows(client, start=None, end=None):
    rates = _in_range(ClientRate.objects.filter(client=client), start, end).select_related(
        "route", "fuel_product", "vehicle_type", "sub_category")
    rows = []
    for r in rates:
        row = [
            r.route.route_code.upper(), (r.fuel_product.name.upper() if r.fuel_product_id else ""),
            r.current_fuel_price, r.current_rate, r.effective_percent, r.rate_subject_to_revision,
            r.updated_fuel_price, r.fuel_price_change_percent, r.rate_adjustment, r.updated_trip_cost,
            r.weight_tons, (r.vehicle_type.name.upper() if r.vehicle_type_id else ""), r.effective_date,
        ]
        if client.has_sub_categories:
            row = [(r.sub_category.name if r.sub_category_id else ""), r.get_rate_type_display().split(" ")[0]] + row
        rows.append(row)
    return rows


def _dedicated_rows(client, start=None, end=None):
    rates = _in_range(DedicatedRate.objects.filter(client=client), start, end).select_related(
        "vehicle", "route", "vehicle_type")
    return [[
        r.vehicle.vehicle_number, r.fixed_cost, r.month.strftime("%b-%y") if r.month else "",
        r.fuel_avg, r.fuel_price, r.variable_cost, (r.route.route_code.upper() if r.route_id else ""),
        r.get_distance_mode_display(), r.distance_km, r.weight_tons,
        (r.vehicle_type.name.upper() if r.vehicle_type_id else ""), r.effective_date,
    ] for r in rates]


def _filename(client, ext):
    return f'{slugify(client.name) or "client"}-rates-{timezone.localdate():%Y%m%d}.{ext}'


@login_required
def client_rates_excel(request, client_id):
    client = get_object_or_404(Client, id=client_id)
    start, end = _date_range(request)
    wb = Workbook()
    bold = Font(bold=True, name="Segoe UI", size=9)
    normal = Font(name="Segoe UI", size=9)
    head_fill = PatternFill("solid", fgColor="F28C28")
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, title, headers, rows):
        ws.cell(row=1, column=1, value=f"{client.name} - {title} ({_period_label(start, end)})").font = Font(bold=True, size=12)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.font = Font(bold=True, name="Segoe UI", size=9, color="FFFFFF")
            cell.fill = head_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r, row in enumerate(rows, start=4):
            for c, value in enumerate(row, start=1):
                if isinstance(value, Decimal):
                    value = float(value)
                cell = ws.cell(row=r, column=c, value=value)
                cell.font = normal
                cell.border = border
                if headers[c - 1] in PERCENT_COLS and isinstance(value, float):
                    cell.number_format = '0.00"%"'
                elif isinstance(value, float):
                    cell.number_format = "#,##0.00"
                elif hasattr(value, "strftime"):
                    cell.number_format = "d-mmm-yy"
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 15
        ws.row_dimensions[3].height = 32
        ws.freeze_panes = "A4"
        if rows:
            ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{len(rows) + 3}"

    ws = wb.active
    ws.title = "Rate Details"
    sheet(ws, "Rate Details", _rate_headers(client), _rate_rows(client, start, end))
    dedicated = _dedicated_rows(client, start, end)
    if dedicated:
        sheet(wb.create_sheet("Dedicated Rates"), "Dedicated Rates", DEDICATED_HEADERS, dedicated)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{_filename(client, "xlsx")}"'
    wb.save(response)
    return response


def _fmt(value, header=""):
    if value is None or value == "":
        return "--"
    if isinstance(value, Decimal):
        return f"{value:,.2f}%" if header in PERCENT_COLS else f"{value:,.2f}"
    if hasattr(value, "strftime"):
        return value.strftime("%d-%b-%y")
    return str(value)


@login_required
def client_rates_pdf(request, client_id):
    client = get_object_or_404(Client, id=client_id)
    start, end = _date_range(request)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=8 * mm, rightMargin=8 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=f"{client.name} - Client Rates")
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    head_style = cell_style.clone("head", fontName="Helvetica-Bold", textColor=colors.white)
    elements = [
        Paragraph(f"{client.name} - Client Rates", styles["Title"]),
        Paragraph(f"{_period_label(start, end)} | Generated {timezone.localdate():%d-%b-%Y}", styles["Normal"]),
        Spacer(1, 4 * mm),
    ]

    def table(title, headers, rows):
        elements.append(Paragraph(title, styles["Heading2"]))
        data = [[Paragraph(h, head_style) for h in headers]]
        data += [[Paragraph(_fmt(v, headers[i]), cell_style) for i, v in enumerate(row)] for row in rows]
        if not rows:
            data.append([Paragraph("No entries", cell_style)] + [""] * (len(headers) - 1))
        width = (297 - 16) * mm / len(headers)
        t = Table(data, colWidths=[width] * len(headers), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f28c28")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9ca3af")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 6 * mm))

    table("Rate Details", _rate_headers(client), _rate_rows(client, start, end))
    dedicated = _dedicated_rows(client, start, end)
    if dedicated:
        table("Dedicated Rates", DEDICATED_HEADERS, dedicated)

    doc.build(elements)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(client, "pdf")}"'
    return response
