import datetime
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from masters.models import City, Route, Client, FuelProduct, VehicleType, ClientRate

CLIENT_NAME = "Five Star 3PL SERVICES"
SOURCE_FILE = "FIVE star rate file aug (upload).xlsx"

# Destination city code -> (full name, latitude, longitude, approx KM from
# Karachi by road). Coordinates/distances are best-effort placeholders where
# the real Route/City record doesn't already exist - get_or_create leaves
# any existing (real) record untouched, only filling gaps for new ones.
CITY_DATA = {
    "KHI": ("KARACHI", Decimal("24.8607"), Decimal("67.0011"), 0),
    "GUJ": ("GUJRANWALA", Decimal("32.1877"), Decimal("74.1945"), 1300),
    "ISB": ("ISLAMABAD", Decimal("33.6844"), Decimal("73.0479"), 1410),
    "LHE": ("LAHORE", Decimal("31.5497"), Decimal("74.3436"), 1225),
    "MUX": ("MULTAN", Decimal("30.1575"), Decimal("71.5249"), 950),
    "PEW": ("PESHAWAR", Decimal("34.0151"), Decimal("71.5249"), 1620),
    "SWL": ("SAHIWAL", Decimal("30.6682"), Decimal("73.1114"), 1100),
    "SKT": ("SIALKOT", Decimal("32.4945"), Decimal("74.5229"), 1370),
    "QTA": ("QUETTA", Decimal("30.1798"), Decimal("66.9750"), 700),
    "SKZ": ("SHIKARPUR", Decimal("27.9560"), Decimal("68.6382"), 500),
    "LRK": ("LARKANA", Decimal("27.5590"), Decimal("68.2120"), 470),
    "SGD": ("SARGODHA", Decimal("32.0836"), Decimal("72.6711"), 1160),
    "BWP": ("BAHAWALPUR", Decimal("29.4000"), Decimal("71.6833"), 750),
    "HYD": ("HYDERABAD", Decimal("25.3960"), Decimal("68.3578"), 165),
    "FSD": ("FAISALABAD", Decimal("31.4180"), Decimal("73.0790"), 1070),
}

# Excel just says "20FT"/"40 FT"/"50FT"/"55 FT" with no DRY/REEFER - client
# confirmed DRY for all sizes.
VEHICLE_SIZE_MAP = {
    "20FT": "20FT DRY",
    "40FT": "40FT DRY",
    "50FT": "50FT DRY",
    "55FT": "55FT DRY",
}


def _dec(value):
    return Decimal(str(value)).quantize(Decimal("0.01"))


class Command(BaseCommand):
    help = f"Import Client Rate entries for '{CLIENT_NAME}' from '{SOURCE_FILE}'."

    def handle(self, *args, **options):
        file_path = Path(SOURCE_FILE)
        if not file_path.exists():
            self.stderr.write(self.style.ERROR(f"'{SOURCE_FILE}' not found in project root."))
            return

        client = Client.objects.filter(name__iexact=CLIENT_NAME).first()
        if not client:
            self.stderr.write(self.style.ERROR(f"Client '{CLIENT_NAME}' not found - create it first."))
            return

        diesel_product, _ = FuelProduct.objects.get_or_create(name="DIESEL")

        khi_name, khi_lat, khi_lng, _ = CITY_DATA["KHI"]
        khi_city, _ = City.objects.get_or_create(
            code="KHI", defaults={"name": khi_name, "latitude": khi_lat, "longitude": khi_lng}
        )

        wb = load_workbook(file_path, data_only=True)
        ws = wb["Sheet1"]

        created, updated, skipped = 0, 0, 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue
            route_code, _product_raw, vtype_raw, weight, eff_date, cur_fuel, cur_rate, eff_pct, upd_fuel = row[:9]

            route_code = str(route_code).strip().upper()
            if "-" not in route_code:
                self.stderr.write(self.style.WARNING(f"Row {idx}: malformed route '{route_code}', skipped."))
                skipped += 1
                continue
            _origin_code, dest_code = route_code.split("-", 1)
            dest_info = CITY_DATA.get(dest_code)
            if not dest_info:
                self.stderr.write(self.style.WARNING(f"Row {idx}: unknown destination city code '{dest_code}', skipped."))
                skipped += 1
                continue

            dest_name, dest_lat, dest_lng, dest_km = dest_info
            dest_city, _ = City.objects.get_or_create(
                code=dest_code, defaults={"name": dest_name, "latitude": dest_lat, "longitude": dest_lng}
            )
            route, _ = Route.objects.get_or_create(
                origin=khi_city, destination=dest_city, defaults={"distance_km": dest_km}
            )

            vtype_key = "".join(str(vtype_raw).strip().upper().split())
            vtype_name = VEHICLE_SIZE_MAP.get(vtype_key)
            if not vtype_name:
                self.stderr.write(self.style.WARNING(f"Row {idx}: unknown vehicle type '{vtype_raw}', skipped."))
                skipped += 1
                continue
            vehicle_type, _ = VehicleType.objects.get_or_create(name=vtype_name)

            eff_date_val = eff_date.date() if isinstance(eff_date, datetime.datetime) else eff_date

            obj, was_created = ClientRate.objects.update_or_create(
                client=client,
                route=route,
                fuel_product=diesel_product,
                vehicle_type=vehicle_type,
                weight_tons=_dec(weight),
                effective_date=eff_date_val,
                defaults={
                    "current_fuel_price": _dec(cur_fuel),
                    "current_rate": _dec(cur_rate),
                    "effective_percent": _dec(eff_pct),
                    "updated_fuel_price": _dec(upd_fuel),
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1

        self.stdout.write(self.style.SUCCESS(
            f"Done. Created {created}, updated {updated}, skipped {skipped} row(s)."
        ))
